import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table


MODULE_PATH = Path(__file__).parents[1] / "scripts" / "actualizar_tbl_normalizado_vaso.py"
SPEC = importlib.util.spec_from_file_location("actualizar_tbl_normalizado_vaso", MODULE_PATH)
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
sys.modules[SPEC.name] = MODULE
SPEC.loader.exec_module(MODULE)


class ClasificacionTests(unittest.TestCase):
    def setUp(self):
        self.rules = [
            MODULE.Rule("Starbucks Con Hielo", "Hel Chai Latte", "Vaso", "2_Helado"),
            MODULE.Rule("Espresso Tradicional", "Latte", "Vaso", "1_Caliente"),
            MODULE.Rule("Espresso Tradicional", "Bundle Contigo", "No", "Na"),
        ]

    def test_hel_se_clasifica_como_bebida_helada(self):
        result = MODULE.suggest("Hel. Pumpkin Latte", "Starbucks Con Hielo", self.rules)
        self.assertTrue(result["auto"])
        self.assertEqual((result["classification"], result["vessel"]), ("Vaso", "2_Helado"))

    def test_latte_sin_marca_fria_se_clasifica_caliente(self):
        result = MODULE.suggest("Pumpkin Latte", "Espresso Tradicional", self.rules)
        self.assertTrue(result["auto"])
        self.assertEqual(result["vessel"], "1_Caliente")

    def test_bundle_se_excluye_del_consumo_de_vaso(self):
        result = MODULE.suggest("Bundle Pumpkin", "Espresso Tradicional", self.rules)
        self.assertTrue(result["auto"])
        self.assertEqual((result["classification"], result["vessel"]), ("No", "Na"))

    def test_frappuccino_abreviado_con_f_se_clasifica_helado(self):
        result = MODULE.suggest("Pumpkin Cream F", "Frappuccino Crema", self.rules)
        self.assertEqual((result["classification"], result["vessel"]), ("Vaso", "2_Helado"))

    def test_latte_h_se_clasifica_helado(self):
        result = MODULE.suggest("NuezPcanaCr LatH", "Starbucks Con Hielo", self.rules)
        self.assertEqual(result["vessel"], "2_Helado")

    def test_cold_foam_estacional_es_adicional(self):
        result = MODULE.suggest("CF Pumpkin CB", "Starbucks Con Hielo", self.rules)
        self.assertEqual((result["classification"], result["vessel"]), ("No", "Na"))

    def test_rtd_no_se_cuenta_como_bebida_preparada(self):
        for name in ("Aranchiatta", "Agua Evian", "Agua Coco"):
            result = MODULE.suggest(name, "Bebidas Frias", self.rules)
            self.assertEqual((result["classification"], result["vessel"]), ("No", "Na"))

    def test_nivel_precio_define_tamano(self):
        self.assertEqual(MODULE.PRICE_SIZE[3], "Grande")
        self.assertEqual(MODULE.PRICE_SIZE[4], "Venti")

    def test_lee_motor_crudo_y_actualiza_tbl_sin_power_query_intermedio(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            motor = root / "motor.xlsx"
            wb = Workbook()
            products = wb.active
            products.title = "producto_base"
            products.append(list(MODULE.PRODUCT_REQUIRED))
            products.append([2715, "Bebidas", "Espresso Tradicional", "Pumpkin Latte"])
            detail = wb.create_sheet("detalleventa_ac")
            detail.append(list(MODULE.RAW_REQUIRED))
            detail.append([2715, 2, 3])
            wb.save(motor)
            rows, sheet, header = MODULE.operational_rows(motor)
            self.assertEqual((sheet, header), ("detalleventa_ac", 1))
            self.assertEqual(rows[0]["description"], "Pumpkin Latte")
            self.assertEqual(rows[0]["quantity"], 2)

            catalog = root / "tbl.xlsx"
            output = root / "tbl_actualizada.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Normalizado_Vaso"
            ws.append(list(MODULE.CATALOG_HEADERS))
            ws.append(["Espresso Tradicional", "Latte", "Vaso", "2_Helado"])
            ws.add_table(Table(displayName="Tabla2", ref="A1:D2"))
            wb.save(catalog)
            changes = [
                {"action": "Correccion", "description": "Latte", "classification": "Vaso", "vessel": "1_Caliente"},
                {"action": "Alta", "family": "Espresso Tradicional", "description": "Pumpkin Latte", "classification": "Vaso", "vessel": "1_Caliente"},
            ]
            MODULE.update_catalog(catalog, output, changes, "Normalizado_Vaso", 1)
            rules, _, _ = MODULE.read_catalog(output)
            self.assertEqual([(r.description, r.vessel) for r in rules], [("Latte", "1_Caliente"), ("Pumpkin Latte", "1_Caliente")])


if __name__ == "__main__":
    unittest.main()
