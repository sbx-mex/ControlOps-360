#!/usr/bin/env python3
"""Verifica las únicas plantillas permitidas, sus macros y el manual de descarga."""
from __future__ import annotations

import argparse
import hashlib
from html.parser import HTMLParser
import json
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
APPROVED_DOWNLOADS = {
    "downloads/motores/Motor_01_Auditoria_Transacciones.xlsm": "7fb5a428d9d57879be47f53e04d0e21e433ef4da284c439b51cbf05414b8bb5e",
    "downloads/motores/Motor_02_Uso_Stock.xlsm": "6929415e6014f4f180c11b2082f66cc1d6fa603009730223cde527c8a553ea08",
    "downloads/motores/Motor_03_Detalle_Productos.xlsm": "7bb2bd9107f509a4ae1a95dd26bb0f390af2644be1857a130e20c53e2000f0b6",
    "downloads/Manual_Control_Ops_360.pdf": "f0a6caa8a2e10150b062f4d4ea3f04ccbd1b821220ba051c1c9ba1b2868565ab",
}
APPROVED_TEMPLATES = frozenset(name for name in APPROVED_DOWNLOADS if name.endswith('.xlsm'))
EXPECTED_TABLES = {
    "01": {"descuento_base", "descuento_ac", "venta_dia_base", "venta_dia_ac", "ticket_base", "ticket_ac", "void_base", "void_ac", "pagos_base", "pagos_ac", "vta_producto_base", "empleado_base", "Tienda"},
    "02": {"uso_base", "uso_ac", "stock_base", "Tienda"},
    "03": {"detalleventa_base", "detalleventa_ac", "producto_base", "Tienda"},
}


class DownloadError(AssertionError):
    pass


class ResourceLinks(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
        self.ids = []

    def handle_starttag(self, tag, attrs):
        values = dict(attrs)
        if values.get('id'):
            self.ids.append(values['id'])
        if tag == 'a' and values.get('href', '').startswith('downloads/'):
            self.links.append(values)


def inspect_template(path: Path, expected_hash: str) -> dict:
    data = path.read_bytes()
    if hashlib.sha256(data).hexdigest() != expected_hash:
        raise DownloadError(f"Plantilla modificada o con información agregada: {path.name}")
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise DownloadError(f"Excel dañado: {path.name}")
        names = archive.namelist()
        if len(names) != len(set(names)):
            raise DownloadError(f"Componentes duplicados: {path.name}")
        for required in ('xl/vbaProject.bin', 'xl/connections.xml', 'customXml/item1.xml'):
            if required not in names:
                raise DownloadError(f"Falta macro o consulta: {path.name} / {required}")
        connections = ET.fromstring(archive.read('xl/connections.xml'))
        if not len(connections):
            raise DownloadError(f"Sin conexiones: {path.name}")
        table_count = sum(name.startswith('xl/tables/table') and name.endswith('.xml') for name in names)
        macro_hash = hashlib.sha256(archive.read('xl/vbaProject.bin')).hexdigest()
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        motor = path.name.split('_')[1]
        expected = EXPECTED_TABLES[motor]
        if set(book.sheetnames) != expected | {'Portada'}:
            raise DownloadError(f"Hojas inesperadas: {path.name}")
        if table_count != len(expected):
            raise DownloadError(f"Tablas incompletas: {path.name}")
        for sheet in book:
            if sheet.title in {'Portada', 'Tienda'}:
                continue
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if any(value not in (None, '') for value in row):
                    raise DownloadError(f"Hay registros operativos: {path.name} / {sheet.title}")
    finally:
        book.close()
    return {'archivo': path.name, 'bytes': len(data), 'tablas': table_count, 'registros_operativos': 0, 'macro_sha256': macro_hash, 'consultas_conservadas': True}


def audit_downloads(root: Path = ROOT) -> dict:
    templates = []
    for name, expected in APPROVED_DOWNLOADS.items():
        path = root / name
        if not path.is_file() or path.is_symlink():
            raise DownloadError(f"Falta descarga: {name}")
        if hashlib.sha256(path.read_bytes()).hexdigest() != expected:
            raise DownloadError(f"Huella de descarga incorrecta: {name}")
        if name in APPROVED_TEMPLATES:
            templates.append(inspect_template(path, expected))
        elif not path.read_bytes().startswith(b'%PDF-'):
            raise DownloadError('Manual PDF no válido')
    parser = ResourceLinks()
    parser.feed((root / 'index.html').read_text(encoding='utf-8'))
    if len(parser.ids) != len(set(parser.ids)):
        raise DownloadError('Identificadores HTML duplicados')
    if len(parser.links) != 4 or {link['href'] for link in parser.links} != APPROVED_DOWNLOADS.keys():
        raise DownloadError('Las descargas visibles no coinciden con los cuatro recursos autorizados')
    for link in parser.links:
        if link.get('download') != Path(link['href']).name:
            raise DownloadError('Nombre de descarga incorrecto')
    return {'estado': 'VERDE', 'plantillas': templates, 'manual_pdf': True, 'enlaces_verificados': len(parser.links)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--json', action='store_true')
    parser.parse_args()
    try:
        print(json.dumps(audit_downloads(), ensure_ascii=False, indent=2))
        return 0
    except Exception as error:
        print(json.dumps({'estado': 'ROJO', 'error': str(error)}, ensure_ascii=False))
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
