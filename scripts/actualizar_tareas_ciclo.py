#!/usr/bin/env python3
"""Extrae el catálogo no operativo de Tareas de Ciclo a un módulo web pequeño."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


CYCLE_COLUMNS = {30: "B", 20: "D", 12: "F", 8: "H"}
CLEANUPS = {
    "utencilios": "utensilios",
    "Utencilios": "Utensilios",
    "Desincrustacion": "Desincrustación",
    "quimicos": "químicos",
    "bascula": "báscula",
    "metalicos": "metálicos",
    "maquina": "máquina",
    "Menus": "menús",
    "estanteria": "estantería",
    "merrycheff": "Merrychef",
    "turbocheff": "Turbochef",
    "Grab&Go": "Grab & Go",
}


def clean_task(value: object) -> str:
    """Normaliza espacios y errores tipográficos sin cambiar el sentido."""

    text = re.sub(r"\s+", " ", str(value or "")).strip().rstrip(".")
    for source, target in CLEANUPS.items():
        text = text.replace(source, target)
    return text


def extract_tasks(workbook_path: Path) -> dict[int, list[str]]:
    """Lee únicamente la hoja de catálogo; nunca exporta datos operativos."""

    workbook = load_workbook(workbook_path, read_only=True, data_only=True, keep_vba=True)
    if "Actividades" not in workbook.sheetnames:
        raise ValueError("El libro no contiene la hoja Actividades.")
    sheet = workbook["Actividades"]
    catalog: dict[int, list[str]] = {}
    for minutes, column in CYCLE_COLUMNS.items():
        tasks: list[str] = []
        seen: set[str] = set()
        for row in range(5, sheet.max_row + 1):
            task = clean_task(sheet[f"{column}{row}"].value)
            key = task.casefold()
            if not task or key in seen:
                continue
            seen.add(key)
            tasks.append(task)
        if not tasks:
            raise ValueError(f"No se encontraron tareas para el ciclo de {minutes} minutos.")
        catalog[minutes] = tasks
    return catalog


def render_module(catalog: dict[int, list[str]]) -> str:
    payload = json.dumps({str(key): value for key, value in catalog.items()}, ensure_ascii=False, indent=2)
    return (
        "// Generado por scripts/actualizar_tareas_ciclo.py; no contiene datos operativos.\n"
        "export const CYCLE_SOURCE='Asistente PH · Tarea de Ciclo';\n"
        f"export const CYCLE_TASKS=Object.freeze({payload});\n"
    )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    catalog = extract_tasks(args.workbook)
    args.output.write_text(render_module(catalog), encoding="utf-8")
    print(json.dumps({"estado": "VERDE", "tareas": sum(map(len, catalog.values())), "salida": str(args.output)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
