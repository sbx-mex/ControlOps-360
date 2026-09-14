#!/usr/bin/env python3
"""Actualiza tbl_normalizado_vaso.xlsx desde un motor Normalizados.

La clasificación exacta siempre tiene prioridad. Las altas automáticas se limitan
a reglas fuertes o coincidencias de alta confianza; los casos dudosos quedan en
un CSV para revisión humana.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import sys
import unicodedata
from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table
except ImportError as error:  # pragma: no cover
    raise SystemExit("Falta openpyxl. Instala con: python -m pip install openpyxl") from error


CATALOG_HEADERS = ("DescripcionFam", "Descripcion", "Normalizado", "Vaso")
OP_REQUIRED = ("DescripcionFam", "Descripcion", "NivelPrecio")
PRICE_SIZE = {1: "Corto", 2: "Alto", 3: "Grande", 7: "Grande", 4: "Venti", 8: "Venti", 9: "Traveler"}
BEVERAGE_FAMILY = re.compile(
    r"espresso|frappuccino|starbucks\s*(con\s*hielo|tea)|alternativas?\s+al\s+cafe|caf[eé]\s+clasico",
    re.I,
)
COLD_NAME = re.compile(
    r"(^|\W)(hel\.?|helado|iced|frozen|frapp|cold\s*brew|refresher|shaken|dragon\s*drink|pink\s*drink)(\W|$)",
    re.I,
)
HOT_NAME = re.compile(
    r"(^|\W)(cal\.?|caliente|cappuccino|cortado|espresso|americano|mocha|latte|chocolate|chai|te\s)(\W|$)",
    re.I,
)
NON_DRINK = re.compile(r"bundle|bndl|contigo|bakery|pastel|crois|pancho|pavpan|salty", re.I)


def normalize(value: object) -> str:
    text = unicodedata.normalize("NFD", str(value or "").strip())
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def tokens(value: object) -> set[str]:
    text = unicodedata.normalize("NFD", str(value or "").lower())
    text = "".join(char for char in text if not unicodedata.combining(char))
    return {part for part in re.findall(r"[a-z0-9]+", text) if len(part) > 1}


def find_header(ws, required: tuple[str, ...], max_rows: int = 25) -> tuple[int, dict[str, int]]:
    if ws.max_row is None or ws.max_column is None:
        ws.calculate_dimension(force=True)
    required_keys = {normalize(value): value for value in required}
    for row_number in range(1, min(max_rows, ws.max_row) + 1):
        values = [ws.cell(row_number, column).value for column in range(1, ws.max_column + 1)]
        positions = {normalize(value): index + 1 for index, value in enumerate(values) if value not in (None, "")}
        if all(key in positions for key in required_keys):
            return row_number, {original: positions[key] for key, original in required_keys.items()}
    raise ValueError(f"{ws.title}: no se encontraron encabezados {', '.join(required)}")


def as_int(value: object) -> int | None:
    try:
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        return None


def as_float(value: object) -> float:
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        return 0.0


@dataclass(frozen=True)
class Rule:
    family: str
    description: str
    normalized: str
    vessel: str


def read_catalog(path: Path) -> tuple[list[Rule], str, int]:
    wb = load_workbook(path, read_only=True, data_only=True)
    found = None
    for ws in wb.worksheets:
        try:
            header_row, columns = find_header(ws, CATALOG_HEADERS)
        except ValueError:
            continue
        found = (ws, header_row, columns)
        break
    if found is None:
        raise ValueError("El catálogo no contiene la estructura de tbl_normalizado_vaso.")
    ws, header_row, columns = found
    rules: list[Rule] = []
    seen: dict[str, tuple[str, str]] = {}
    for row_number in range(header_row + 1, ws.max_row + 1):
        description = str(ws.cell(row_number, columns["Descripcion"]).value or "").strip()
        if not description:
            continue
        family = str(ws.cell(row_number, columns["DescripcionFam"]).value or "").strip()
        classification = str(ws.cell(row_number, columns["Normalizado"]).value or "").strip()
        vessel = str(ws.cell(row_number, columns["Vaso"]).value or "").strip()
        if classification not in {"Vaso", "No", "FHW"}:
            raise ValueError(f"Clasificación inválida en {ws.title}!{row_number}: {classification}")
        if (classification == "Vaso" and vessel not in {"1_Caliente", "2_Helado"}) or (
            classification == "FHW" and vessel != "3_FHW"
        ) or (classification == "No" and vessel != "Na"):
            raise ValueError(f"Regla contradictoria en {ws.title}!{row_number}: {description}")
        key = normalize(description)
        value = (classification, vessel)
        if key in seen and seen[key] != value:
            raise ValueError(f"Cruce ambiguo en catálogo: {description}")
        seen[key] = value
        rules.append(Rule(family, description, classification, vessel))
    return rules, ws.title, header_row


def operational_rows(path: Path) -> tuple[list[dict[str, object]], str, int]:
    wb = load_workbook(path, read_only=True, data_only=True, keep_vba=path.suffix.lower() == ".xlsm")
    candidates = []
    for preferred in ("normalizados_ac", "normalizados_base"):
        ws = next((sheet for sheet in wb.worksheets if normalize(sheet.title) == normalize(preferred)), None)
        if ws is None:
            continue
        header_row, columns = find_header(ws, OP_REQUIRED)
        optional = {}
        values = [ws.cell(header_row, column).value for column in range(1, ws.max_column + 1)]
        normalized_headers = {normalize(value): index + 1 for index, value in enumerate(values) if value not in (None, "")}
        for name in ("CodigoDIA", "CantidadAjustada", "IDTamanio"):
            optional[name] = normalized_headers.get(normalize(name))
        rows = []
        for row_number in range(header_row + 1, ws.max_row + 1):
            description = str(ws.cell(row_number, columns["Descripcion"]).value or "").strip()
            if not description:
                continue
            quantity = as_float(ws.cell(row_number, optional["CantidadAjustada"]).value) if optional["CantidadAjustada"] else 1.0
            if quantity <= 0:
                continue
            rows.append(
                {
                    "family": str(ws.cell(row_number, columns["DescripcionFam"]).value or "").strip(),
                    "description": description,
                    "price_level": as_int(ws.cell(row_number, columns["NivelPrecio"]).value),
                    "dia": str(ws.cell(row_number, optional["CodigoDIA"]).value or "").strip() if optional["CodigoDIA"] else "",
                    "quantity": quantity,
                }
            )
        candidates.append((rows, ws.title, header_row))
        if rows:
            return rows, ws.title, header_row
    return candidates[0] if candidates else ([], "", 0)


def similarity(description: str, family: str, rule: Rule) -> float:
    left, right = normalize(description), normalize(rule.description)
    sequence = SequenceMatcher(None, left, right).ratio()
    lt, rt = tokens(description), tokens(rule.description)
    overlap = len(lt & rt) / len(lt | rt) if lt | rt else 0.0
    family_bonus = 0.08 if normalize(family) == normalize(rule.family) else 0.0
    return min(1.0, sequence * 0.78 + overlap * 0.22 + family_bonus)


def suggest(description: str, family: str, rules: list[Rule]) -> dict[str, object]:
    ranked = sorted(((similarity(description, family, rule), rule) for rule in rules), key=lambda item: item[0], reverse=True)
    best_score, best = ranked[0]
    second_score = ranked[1][0] if len(ranked) > 1 else 0.0
    margin = best_score - second_score
    cold = bool(COLD_NAME.search(description)) or bool(re.search(r"frappuccino|con\s+hielo", family, re.I))
    hot = not cold and (bool(HOT_NAME.search(description)) or bool(re.search(r"cafe\s+clasico", family, re.I)))
    excluded = bool(NON_DRINK.search(description))
    beverage = bool(BEVERAGE_FAMILY.search(family)) or cold or hot

    classification, vessel, method, confidence = "", "", "Sin regla suficiente", best_score
    auto = False
    if excluded:
        method = "Nombre de paquete o complemento; requiere revisión"
    elif beverage and cold:
        classification, vessel, method, confidence, auto = "Vaso", "2_Helado", "Regla explícita de bebida fría", max(best_score, 0.96), True
    elif beverage and hot:
        classification, vessel, method, confidence, auto = "Vaso", "1_Caliente", "Regla explícita de bebida caliente", max(best_score, 0.94), True
    elif best_score >= 0.92 and margin >= 0.05 and normalize(family) == normalize(best.family):
        classification, vessel = best.normalized, best.vessel
        method, confidence, auto = "Coincidencia alta en la misma familia", best_score, True
    elif best_score >= 0.80:
        classification, vessel, method = best.normalized, best.vessel, "Coincidencia para revisión"

    return {
        "classification": classification,
        "vessel": vessel,
        "method": method,
        "confidence": confidence,
        "auto": auto and classification in {"Vaso", "No", "FHW"},
        "similar": best.description,
        "similar_family": best.family,
        "margin": margin,
    }


def update_catalog(source: Path, destination: Path, additions: list[dict[str, object]], sheet_name: str, header_row: int) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    if not additions:
        shutil.copy2(source, destination)
        return
    wb = load_workbook(source)
    ws = wb[sheet_name]
    _, columns = find_header(ws, CATALOG_HEADERS)
    start_row = ws.max_row + 1
    style_row = ws.max_row
    for offset, item in enumerate(additions):
        row_number = start_row + offset
        values = (item["family"], item["description"], item["classification"], item["vessel"])
        for header, value in zip(CATALOG_HEADERS, values):
            target = ws.cell(row_number, columns[header], value)
            source_cell = ws.cell(style_row, columns[header])
            if source_cell.has_style:
                target._style = copy(source_cell._style)
            target.number_format = source_cell.number_format
            target.alignment = copy(source_cell.alignment)
        ws.row_dimensions[row_number].height = ws.row_dimensions[style_row].height
    last_row = ws.max_row
    for table in ws.tables.values():
        if isinstance(table, Table) and table.ref:
            min_ref, max_ref = table.ref.split(":")
            if re.match(rf"^[A-Z]+{header_row}$", min_ref):
                table.ref = f"{min_ref}:{get_column_letter(ws.max_column)}{last_row}"
    wb.save(destination)


def main() -> int:
    parser = argparse.ArgumentParser(description="Detecta y propone bebidas faltantes en tbl_normalizado_vaso.")
    parser.add_argument("--operativo", required=True, type=Path, help="Motor_Normalizados actualizado (.xlsm/.xlsx)")
    parser.add_argument("--catalogo", required=True, type=Path, help="tbl_normalizado_vaso.xlsx vigente")
    parser.add_argument("--salida", required=True, type=Path, help="Copia actualizada del tbl")
    parser.add_argument("--reporte", required=True, type=Path, help="CSV de altas y pendientes")
    parser.add_argument("--solo-reporte", action="store_true", help="No agrega ni siquiera coincidencias de alta confianza")
    args = parser.parse_args()

    rules, sheet_name, header_row = read_catalog(args.catalogo)
    sales, source_sheet, source_header = operational_rows(args.operativo)
    existing = {normalize(rule.description) for rule in rules}
    grouped: dict[str, dict[str, object]] = {}
    for row in sales:
        key = normalize(row["description"])
        if key in existing:
            continue
        item = grouped.setdefault(
            key,
            {"family": row["family"], "description": row["description"], "quantity": 0.0, "levels": Counter(), "dia": Counter()},
        )
        item["quantity"] += float(row["quantity"])
        if row["price_level"] is not None:
            item["levels"][row["price_level"]] += float(row["quantity"])
        if row["dia"]:
            item["dia"][row["dia"]] += float(row["quantity"])

    report_rows = []
    additions = []
    for item in sorted(grouped.values(), key=lambda value: (-value["quantity"], str(value["description"]))):
        proposal = suggest(str(item["description"]), str(item["family"]), rules)
        levels = [level for level, _ in item["levels"].most_common()]
        sizes = []
        for level in levels:
            size = PRICE_SIZE.get(level, f"Nivel {level} sin mapa")
            if size not in sizes:
                sizes.append(size)
        dominant_level = levels[0] if levels else None
        row = {
            **item,
            **proposal,
            "dia_value": item["dia"].most_common(1)[0][0] if item["dia"] else "",
            "levels_value": ", ".join(map(str, levels)),
            "size": PRICE_SIZE.get(dominant_level, "Sin tamaño") if dominant_level is not None else "Sin tamaño",
            "sizes_value": ", ".join(sizes),
        }
        report_rows.append(row)
        if proposal["auto"] and not args.solo_reporte:
            additions.append(row)

    args.reporte.parent.mkdir(parents=True, exist_ok=True)
    fields = [
        "DescripcionFam", "Descripcion", "CodigoDIA", "Cantidad", "NivelPrecio", "Tamaño dominante", "Tamaños observados",
        "Normalizado propuesto", "Vaso propuesto", "Confianza", "Método", "Similar a", "Alta automática",
    ]
    with args.reporte.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=fields)
        writer.writeheader()
        for row in report_rows:
            writer.writerow(
                {
                    "DescripcionFam": row["family"],
                    "Descripcion": row["description"],
                    "CodigoDIA": row["dia_value"],
                    "Cantidad": round(row["quantity"], 3),
                    "NivelPrecio": row["levels_value"],
                    "Tamaño dominante": row["size"],
                    "Tamaños observados": row["sizes_value"],
                    "Normalizado propuesto": row["classification"],
                    "Vaso propuesto": row["vessel"],
                    "Confianza": f"{row['confidence']:.1%}",
                    "Método": row["method"],
                    "Similar a": row["similar"],
                    "Alta automática": "Si" if row["auto"] and not args.solo_reporte else "No",
                }
            )

    update_catalog(args.catalogo, args.salida, additions, sheet_name, header_row)
    print(
        json.dumps(
            {
                "operativo": str(args.operativo),
                "hoja_operativa": source_sheet,
                "fila_encabezado_operativa": source_header,
                "filas_operativas": len(sales),
                "bebidas_sin_regla": len(report_rows),
                "altas_automaticas": len(additions),
                "catalogo_salida": str(args.salida),
                "reporte": str(args.reporte),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (OSError, ValueError) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(2)
