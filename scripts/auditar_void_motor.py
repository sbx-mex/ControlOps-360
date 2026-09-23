#!/usr/bin/env python3
"""Reconciliación local de cheques Void en motores externos, sin modificar XLSM."""

from __future__ import annotations

import argparse
import json
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from compatibilidad_xlsm import inspect_workbook


def day(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, (float, int)):
        return (datetime(1899, 12, 30) + timedelta(days=value)).date().isoformat()
    text = str(value).strip()
    return datetime.fromisoformat(text.replace("Z", "")).date().isoformat()


def records(book, source):
    min_col, min_row, max_col, max_row = range_boundaries(source.reference)
    rows = book[source.sheet].iter_rows(
        min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=True
    )
    headers = list(source.columns)
    next(rows)
    for cells in rows:
        if cells[0] is not None:
            yield dict(zip(headers, cells))


def key(row, date_column):
    return str(row["IDTienda"]).removesuffix(".0"), day(row[date_column]), str(row["Ticket"]).removesuffix(".0")


def audit(motor_one: Path, motor_two: Path | None = None):
    first = inspect_workbook(motor_one)
    if not first.compatible:
        raise ValueError(f"Motor 1 incompatible: {first.reason}")
    sources = {source.sheet: source for source in first.sources}
    required = {"ticket_ac", "void_ac", "pagos_ac"}
    if not required.issubset(sources):
        raise ValueError(f"Motor 1 no incluye {', '.join(sorted(required - sources.keys()))}")
    if motor_two:
        second = inspect_workbook(motor_two)
        if not second.compatible or not {"uso_ac", "stock_base"}.issubset({s.sheet for s in second.sources}):
            raise ValueError("Motor 2 no tiene uso_ac y stock_base compatibles")
        if set(first.cecos) != set(second.cecos):
            raise ValueError("Los motores no corresponden al mismo CeCo")
    book = load_workbook(motor_one, read_only=True, data_only=True)
    try:
        tickets = {key(r, "FechaCierre"): float(r["Total"]) for r in records(book, sources["ticket_ac"])}
        payments = defaultdict(float)
        for row in records(book, sources["pagos_ac"]):
            payments[key(row, "FechaHora")] += float(row["MontoTotal"] or 0)
        checks = defaultdict(lambda: {"lines": 0, "negative": 0, "reopen": False})
        for row in records(book, sources["void_ac"]):
            check = checks[key(row, "FechaHora")]
            check["lines"] += 1
            check["negative"] += int(float(row["Total"] or 0) < 0)
            check["reopen"] |= str(row["VoidReason"] or "").strip().lower().startswith("r")
    finally:
        book.close()
    counts = defaultdict(int)
    for identity, check in checks.items():
        total = tickets.get(identity)
        payment = payments.get(identity)
        if total is not None and total < 0:
            category = "void_exclusivo"
        elif check["negative"] and ((total is not None and total > 0) or (total is None and payment is not None and payment > 0)):
            category = "riesgo_borrado"
            if total is None:
                counts["borrado_inferido_por_pagos"] += 1
        else:
            category = "sin_cierre_verificado"
        counts[category] += 1
        counts["ticket_reabierto"] += int(check["reopen"])
        counts["lineas_void"] += check["lines"]
        counts["lineas_negativas"] += check["negative"]
    return {"cecos": first.cecos, "cheques": len(checks), **dict(sorted(counts.items()))}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("motor_01", type=Path)
    parser.add_argument("motor_02", type=Path, nargs="?")
    args = parser.parse_args()
    print(json.dumps(audit(args.motor_01, args.motor_02), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
