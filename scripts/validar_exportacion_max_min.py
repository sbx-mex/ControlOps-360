#!/usr/bin/env python3
"""Valida estructura y cálculos de una exportación Max & Min."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


ORDER_FACTOR = {2: 5, 3: 4, 4: 3, 5: 2}
EXPORT_HEADERS = (
    "Descripción SAP", "Nombre Micros", "#DIA", "#SAP", "Min", "Max",
    "Unidad / Pick Pack", "Pz / Caja", "# Pedido",
)
META_LABELS = ("TIENDA", "PERIODO INI - FIN", "ACTUALIZACIÓN / IMPRESIÓN", "# PEDIDOS")


def table_rows(sheet, expected_headers: tuple[str, ...]) -> list[dict[str, object]]:
    values = list(sheet.iter_rows(values_only=True))
    if len(values) < 3 or tuple(values[2][:len(expected_headers)]) != expected_headers:
        raise ValueError(f"{sheet.title}: encabezados inesperados")
    metadata = values[0]
    if tuple(metadata[index] for index in (0, 2, 4, 6)) != META_LABELS:
        raise ValueError(f"{sheet.title}: encabezado operativo incompleto")
    if any(metadata[index] in (None, "") for index in (1, 3, 5, 7)):
        raise ValueError(f"{sheet.title}: metadatos operativos incompletos")
    if sheet.freeze_panes != "A4" or sheet.print_title_rows != "$1:$3":
        raise ValueError(f"{sheet.title}: filas superiores no están fijas/repetidas")
    if sheet.sheet_properties.pageSetUpPr.fitToPage is not True or sheet.page_setup.orientation != "landscape" or sheet.page_setup.fitToWidth != 1:
        raise ValueError(f"{sheet.title}: configuración de impresión inesperada")
    return [dict(zip(expected_headers, row)) for row in values[3:] if any(value not in (None, "") for value in row)]


def validate(path: Path) -> dict[str, int]:
    workbook = load_workbook(path, read_only=False, data_only=False)
    if workbook.sheetnames != ["Uso Unidad", "Pick Pack"]:
        raise ValueError(f"Pestañas inesperadas: {workbook.sheetnames}")

    unit_rows = table_rows(workbook["Uso Unidad"], EXPORT_HEADERS)
    pack_rows = table_rows(workbook["Pick Pack"], EXPORT_HEADERS)
    identities = [(row["Descripción SAP"], row["#DIA"], row["#SAP"]) for row in unit_rows]
    if len(identities) != len(set(identities)):
        raise ValueError("Uso Unidad: hay artículos duplicados")

    unit_by_identity = {identity: row for identity, row in zip(identities, unit_rows)}
    for row in unit_rows:
        orders = int(row["# Pedido"])
        minimum = float(row["Min"])
        maximum = float(row["Max"])
        if orders not in ORDER_FACTOR or abs(maximum - minimum * ORDER_FACTOR[orders]) > 0.31:
            raise ValueError(f"Uso Unidad: relación Max & Min inválida en {row['Descripción SAP']}")
        if row["Unidad / Pick Pack"] != "Unidad":
            raise ValueError("Uso Unidad: formato operativo inesperado")
        if row["Pz / Caja"] not in (None, ""):
            raise ValueError("Uso Unidad: no debe mostrar conversión de caja")

    sleeve_rows = 0
    for row in pack_rows:
        identity = (row["Descripción SAP"], row["#DIA"], row["#SAP"])
        if identity not in unit_by_identity:
            raise ValueError("Pick Pack: artículo sin correspondencia en Uso Unidad")
        if row["Unidad / Pick Pack"] not in {"Pick Pack", "Manga"}:
            raise ValueError(f"Pick Pack: formato desconocido {row['Unidad / Pick Pack']}")
        content = int(row["Pz / Caja"])
        if content <= 1:
            raise ValueError(f"Pick Pack: conversión redundante de {content} pieza")
        if row["Min"] not in (None, "") and row["Max"] not in (None, ""):
            orders = int(row["# Pedido"])
            minimum, maximum = int(row["Min"]), int(row["Max"])
            if orders not in ORDER_FACTOR or maximum < minimum or maximum > minimum * ORDER_FACTOR[orders]:
                raise ValueError(f"Pick Pack: relación Max & Min inválida en {row['Descripción SAP']}")
        if row["Unidad / Pick Pack"] == "Manga":
            sleeve_rows += 1
            if content not in {40, 50, 100}:
                raise ValueError(f"Manga inválida: {content} piezas")

    return {"uso_unidad": len(unit_rows), "pick_pack": len(pack_rows), "mangas": sleeve_rows}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("archivo", type=Path)
    args = parser.parse_args()
    print(json.dumps(validate(args.archivo), ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
